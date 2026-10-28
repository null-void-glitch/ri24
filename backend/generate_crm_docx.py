# -*- coding: utf-8 -*-
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DATA = Path(__file__).with_name("crm_features_data.json")
OUT  = Path(__file__).resolve().parent.parent / "Erweiterungen CRM" / "CRM-Erweiterungsvorschlaege.docx"

RI24_BLUE  = RGBColor(0x00, 0x5B, 0x9A)
RI24_LIGHT = RGBColor(0xE8, 0xF2, 0xFA)
SUBTLE     = RGBColor(0x6C, 0x75, 0x7D)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK       = RGBColor(0x1A, 0x1A, 0x2E)
LIGHT_BLUE = RGBColor(0xA8, 0xD8, 0xFF)

def shd(cell, rgb):
    tc = cell._tc
    pr = tc.get_or_add_tcPr()
    s  = OxmlElement("w:shd")
    s.set(qn("w:val"),   "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"),  str(rgb))
    pr.append(s)

def run(para, text, bold=False, size=10.5, color=None):
    r = para.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return r

def main():
    dat = json.loads(DATA.read_text(encoding="utf-8"))
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.5)

    # Defaults
    doc.styles["Normal"].font.name = "Calibri"

    # Title
    p = doc.add_paragraph()
    run(p, "CRM-Erweiterungsvorschlaege", bold=True, size=26, color=RI24_BLUE)
    p = doc.add_paragraph()
    run(p, "Partnerbereich Baufinanzierung  |  Regionalimmobilien24 GmbH  |  Oktober 2026", size=11, color=SUBTLE)
    doc.add_paragraph()

    p = doc.add_paragraph()
    run(p, (
        "Auf Basis einer vollstaendigen automatisierten Analyse des Partner-CRM (/be-admin/) wurden "
        "die folgenden 13 Erweiterungsvorschlaege erarbeitet. Die Analyse umfasste: Kundenverwaltung, "
        "Europace-Integration, Profil, Landingpages, Empfehlungsseite, Unterlagen sowie bestehende "
        "API- und Datenbankfelder. Die Vorschlaege sind nach Umsetzungsaufwand und strategischem Hebel "
        "priorisiert. Die drei staerksten Quick-Wins sind #2, #1 und #13."
    ), size=10.5)
    doc.add_paragraph()

    for feat in dat["features"]:
        # Blue header
        t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
        c = t.rows[0].cells[0]; shd(c, RI24_BLUE)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after  = Pt(6)
        run(p, f"#{feat['num']}  ", bold=True, size=13, color=LIGHT_BLUE)
        run(p, feat["title"],       bold=True, size=13, color=WHITE)

        # Meta table
        mt = doc.add_table(rows=1, cols=4); mt.style = "Table Grid"
        for i, (lbl, val) in enumerate(zip(
            ["Bereich","Aufwand","Hebel","Empfehlung"],
            [feat["area"],feat["effort"],feat["lever"],feat["sprint"]]
        )):
            cell = mt.rows[0].cells[i]; shd(cell, RI24_LIGHT)
            p2 = cell.paragraphs[0]
            p2.paragraph_format.space_before = Pt(4)
            p2.paragraph_format.space_after  = Pt(4)
            run(p2, lbl+"\n", bold=True, size=8,   color=SUBTLE)
            run(p2, val,       bold=True, size=9.5, color=RI24_BLUE)

        doc.add_paragraph()

        for label, content, is_list in [
            ("Ist-Zustand", feat["current"],  False),
            ("Vorschlag",   feat["proposal"], False),
            ("Nutzen",      feat["benefit"],  True),
        ]:
            p = doc.add_paragraph()
            run(p, label, bold=True, size=10, color=RI24_BLUE)
            p.paragraph_format.space_after = Pt(2)
            if is_list:
                for item in content:
                    bp = doc.add_paragraph(style="List Bullet")
                    run(bp, item, size=10.5)
                    bp.paragraph_format.space_after = Pt(2)
            else:
                bp = doc.add_paragraph()
                run(bp, content, size=10.5)
                bp.paragraph_format.space_after = Pt(6)

        doc.add_paragraph()

    # Priority matrix
    doc.add_page_break()
    p = doc.add_paragraph()
    run(p, "Prioritaetsmatrix", bold=True, size=20, color=RI24_BLUE)

    tbl = doc.add_table(rows=1, cols=5); tbl.style = "Table Grid"
    for i, hdr in enumerate(["#","Feature","Aufwand","Hebel","Empfehlung"]):
        cell = tbl.rows[0].cells[i]; shd(cell, RI24_BLUE)
        p2 = cell.paragraphs[0]
        p2.paragraph_format.space_before = p2.paragraph_format.space_after = Pt(4)
        run(p2, hdr, bold=True, size=9.5, color=WHITE)

    for idx, row_data in enumerate(dat["priority"]):
        num, name, effort, lever, sprint = row_data
        row = tbl.add_row()
        bg  = RI24_LIGHT if idx % 2 == 0 else WHITE
        for i, val in enumerate([str(num), name, effort, lever, sprint]):
            cell = row.cells[i]; shd(cell, bg)
            p2 = cell.paragraphs[0]
            p2.paragraph_format.space_before = p2.paragraph_format.space_after = Pt(3)
            bold = (sprint == "Sprint 1") and (i > 0)
            color = RI24_BLUE if sprint == "Sprint 1" and i > 0 else DARK
            run(p2, val, bold=bold, size=9.5, color=color)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "Erstellt auf Basis einer automatisierten CRM-Analyse  |  Regionalimmobilien24 GmbH  |  Oktober 2026",
        size=9, color=SUBTLE)

    doc.save(str(OUT))
    print(f"Saved: {OUT}")

if __name__ == "__main__":
    main()
